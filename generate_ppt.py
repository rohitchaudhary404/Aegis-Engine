import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def add_title(slide, text, color=RGBColor(0, 200, 255)):
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(1))
    tf = title_box.text_frame
    p = tf.add_paragraph()
    p.text = text
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = color
    return tf

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    dark_bg = RGBColor(15, 20, 35)
    text_light = RGBColor(220, 220, 220)
    accent_blue = RGBColor(0, 200, 255)
    accent_green = RGBColor(50, 200, 100)
    accent_red = RGBColor(255, 80, 80)
    
    # ---------------------------------------------------------
    # SLIDE 1: Title Slide
    # ---------------------------------------------------------
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = dark_bg
    
    title_box = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11), Inches(1.5))
    p = title_box.text_frame.add_paragraph()
    p.text = "INTEGRAMINDS"
    p.font.size = Pt(72)
    p.font.bold = True
    p.font.color.rgb = accent_blue
    p.alignment = PP_ALIGN.CENTER
    
    sub_box = slide.shapes.add_textbox(Inches(1), Inches(4.0), Inches(11), Inches(1))
    p2 = sub_box.text_frame.add_paragraph()
    p2.text = "Autonomous, Self-Healing Cybersecurity Assessment Engine"
    p2.font.size = Pt(28)
    p2.font.color.rgb = text_light
    p2.alignment = PP_ALIGN.CENTER

    # ---------------------------------------------------------
    # SLIDE 2: The Problem
    # ---------------------------------------------------------
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = dark_bg
    
    add_title(slide, "The Problem: The 'Slot 1' Vulnerability", accent_red)
    
    tx_box = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(11), Inches(5))
    tf = tx_box.text_frame
    tf.word_wrap = True
    
    p = tf.add_paragraph()
    p.text = "Multi-day recruitment drives (TCS, Infosys, Universities) are structurally broken:"
    p.font.size = Pt(28)
    p.font.color.rgb = text_light
    p.font.bold = True
    
    points = [
        "1. Slot 1 candidates memorize or photograph exams.",
        "2. Questions are immediately leaked to Telegram and WhatsApp.",
        "3. Later slots see unnatural 95%+ passing rates.",
        "4. Traditional testing platforms are passive and cannot react to live leaks."
    ]
    for pt in points:
        p = tf.add_paragraph()
        p.text = pt
        p.font.size = Pt(24)
        p.font.color.rgb = text_light
        p.level = 1

    # ---------------------------------------------------------
    # SLIDE 3: System Architecture (Diagram)
    # ---------------------------------------------------------
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = dark_bg
    add_title(slide, "System Architecture & Flow", accent_blue)
    
    def draw_box(x, y, w, h, text, color):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.color.rgb = RGBColor(255, 255, 255)
        tf = shape.text_frame
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(16)
        p.font.bold = True
        p.alignment = PP_ALIGN.CENTER
        return shape
        
    def draw_arrow(x, y, w, h):
        shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = text_light
        shape.line.color.rgb = text_light
    
    # Diagram Blocks
    draw_box(0.5, 2, 2.5, 1.5, "1. INGESTION\n(PyMuPDF / Regex)\nMessy PDFs -> Clean DB", RGBColor(50, 50, 80))
    draw_arrow(3.1, 2.5, 0.8, 0.5)
    
    draw_box(4.0, 2, 2.5, 1.5, "2. OPTIMIZATION\nMathematical Overlap\nConstraint Engine", RGBColor(50, 80, 50))
    draw_arrow(6.6, 2.5, 0.8, 0.5)
    
    draw_box(7.5, 2, 2.5, 1.5, "3. SECURITY SHIELD\nCUSUM Anomaly Tracking\n& OSINT Dorking", RGBColor(120, 40, 40))
    draw_arrow(10.1, 2.5, 0.8, 0.5)
    
    draw_box(11.0, 2, 2.0, 4.0, "OUTPUT\n\nExported PDFs\n+\nQR Hash\n+\nCanary Tokens", RGBColor(40, 40, 40))
    
    # Feedback loop
    draw_box(4.0, 4.5, 6.0, 1.5, "4. SELF-HEALING AI (Groq LPU)\nIntercepts leak -> Auto-rewrites methodology -> Invalidates Answer Keys", RGBColor(80, 40, 80))
    # Loop arrows
    arr1 = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(8.5), Inches(3.6), Inches(0.5), Inches(0.8))
    arr1.fill.solid()
    arr1.fill.fore_color.rgb = accent_red
    arr1.line.color.rgb = accent_red

    arr2 = slide.shapes.add_shape(MSO_SHAPE.UP_ARROW, Inches(5.0), Inches(3.6), Inches(0.5), Inches(0.8))
    arr2.fill.solid()
    arr2.fill.fore_color.rgb = accent_green
    arr2.line.color.rgb = accent_green
    

    # ---------------------------------------------------------
    # SLIDE 4: The "Masterpiece" Features
    # ---------------------------------------------------------
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = dark_bg
    add_title(slide, "The 'Masterpiece' Features", accent_green)
    
    tx_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(12), Inches(5))
    tf = tx_box.text_frame
    tf.word_wrap = True
    
    features = [
        ("Zero-Shot Autonomous Expansion", "If the test bank is mathematically too small for constraints, Groq AI synthesizes brand new questions in real-time to fill the deficit."),
        ("Forensic Traitor Tracing (Canary Tokens)", "Zero-width invisible cryptographic characters are embedded in the text. If a student copy-pastes a question, our decoder traces it to the exact slot."),
        ("Physical Zero-Trust Authentication", "Printed PDFs receive an SHA-256 Hash and QR code. Live OpenCV webcam scanning authenticates the physical paper in the exam hall."),
        ("Live Threat Topology", "Real-time DuckDuckGo API (site:t.me) finds live leaks and plots the blast radius across domains using a NetworkX topological graph.")
    ]
    
    for title, desc in features:
        p = tf.add_paragraph()
        p.text = f"- {title}: "
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = accent_green
        
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(18)
        p2.font.color.rgb = text_light
        p2.level = 1

    # ---------------------------------------------------------
    # SLIDE 5: Technology Stack
    # ---------------------------------------------------------
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = dark_bg
    add_title(slide, "Technology Stack", accent_blue)
    
    tx_box = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(11), Inches(5))
    tf = tx_box.text_frame
    
    stacks = [
        "Frontend & UI: Streamlit, Plotly (Heatmaps/ICC)",
        "Backend & DB: Python 3.11, SQLite3 (WAL Concurrency)",
        "AI Engine: Groq LPU (openai/gpt-oss-20b), Pydantic (Structured JSON)",
        "Cyber & Math: SymPy (Algebra), CUSUM (Stats), Hashlib (Crypto)",
        "OSINT & Vision: ddgs (DuckDuckGo API), NetworkX, OpenCV (QR)",
        "Doc Processing: PyMuPDF (fitz), Regex, FPDF2"
    ]
    
    for s in stacks:
        p = tf.add_paragraph()
        p.text = s
        p.font.size = Pt(24)
        p.font.color.rgb = text_light
        p.level = 0
        p.space_after = Pt(14)

    # ---------------------------------------------------------
    # SLIDE 6: Conclusion
    # ---------------------------------------------------------
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = dark_bg
    
    title_box = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11), Inches(1.5))
    p = title_box.text_frame.add_paragraph()
    p.text = "Thank You"
    p.font.size = Pt(65)
    p.font.bold = True
    p.font.color.rgb = accent_blue
    p.alignment = PP_ALIGN.CENTER
    
    sub_box = slide.shapes.add_textbox(Inches(1), Inches(4.0), Inches(11), Inches(1))
    p2 = sub_box.text_frame.add_paragraph()
    p2.text = "Ready for the Live Product Demo."
    p2.font.size = Pt(28)
    p2.font.color.rgb = text_light
    p2.alignment = PP_ALIGN.CENTER

    prs.save("Integraminds_Assessment_Engine.pptx")

if __name__ == "__main__":
    create_presentation()



